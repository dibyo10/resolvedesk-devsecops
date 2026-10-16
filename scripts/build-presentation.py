import json
from pathlib import Path
from textwrap import wrap

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from reportlab.lib.colors import HexColor
from reportlab.pdfgen.canvas import Canvas

ROOT = Path(__file__).resolve().parents[1]
NAVY = "102A43"
TEAL = "087F8C"
WHITE = "FFFFFF"
GRAY = "536879"


def build():
    slides = json.loads((ROOT / "demo/slides.json").read_text())
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    pdf = Canvas(str(ROOT / "demo/slides.pdf"), pagesize=(960, 540))
    for index, data in enumerate(slides, 1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(WHITE)
        pdf.setFillColor(HexColor("#" + WHITE))
        pdf.rect(0, 0, 960, 540, fill=1, stroke=0)

        def text(value, x, y, width, height, size=22, color=NAVY, bold=False):
            box = slide.shapes.add_textbox(
                Inches(x), Inches(y), Inches(width), Inches(height)
            )
            frame = box.text_frame
            frame.word_wrap = True
            paragraph = frame.paragraphs[0]
            paragraph.text = value
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.name = "Aptos"
            paragraph.font.color.rgb = RGBColor.from_string(color)
            pdf.setFillColor(HexColor("#" + color))
            pdf.setFont("Helvetica-Bold" if bold else "Helvetica", size)
            for line_number, line in enumerate(
                wrap(
                    value.replace("→", "->").replace("←", "<-"),
                    max(1, int(width * 72 / (size * 0.51))),
                )
            ):
                pdf.drawString(
                    x * 72 + 7, 540 - y * 72 - size - line_number * size * 1.25, line
                )

        text(data["subtitle"].upper(), 0.65, 0.45, 12, 0.45, 12, TEAL, True)
        text(data["title"], 0.65, 1.0, 12, 1.0, 34, NAVY, True)
        start = 2.25
        if data.get("diagram"):
            for x, label in [
                (0.7, "Browser"),
                (3.2, "React / Nginx"),
                (6.0, "FastAPI"),
                (8.8, "PostgreSQL"),
            ]:
                text(label, x, 2.4, 2.5, 0.7, 22, TEAL, True)
            text("→", 2.6, 2.4, 0.5, 0.7)
            text("→", 5.4, 2.4, 0.5, 0.7)
            text("→", 8.2, 2.4, 0.5, 0.7)
            text("GitHub Actions → GHCR → ArgoCD / Helm → EKS", 0.7, 3.4, 12, 0.6, 24)
            text(
                "Prometheus ← /metrics       Grafana ← Prometheus",
                0.7,
                4.2,
                12,
                0.6,
                22,
            )
            start = 5.25
        for point_index, point in enumerate(data["points"]):
            text(
                f"{point_index + 1:02d}",
                0.7,
                start + point_index * 0.79,
                0.5,
                0.5,
                13,
                TEAL,
                True,
            )
            text(point, 1.3, start + point_index * 0.79, 11.1, 0.79, 20)
        text("RESOLVEDESK  /  DEVSECOPS PROJECT", 0.7, 7.0, 10, 0.3, 10, GRAY)
        text(f"{index:02d}", 12.0, 7.0, 0.6, 0.3, 10, GRAY)
        pdf.showPage()
    deck.save(ROOT / "demo/presentation.pptx")
    pdf.save()
    print(f"Built {len(slides)} slides as PowerPoint and PDF")


if __name__ == "__main__":
    build()
